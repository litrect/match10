from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
D=Document(); sec=D.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
 s=D.styles[name]; s.font.name='Microsoft YaHei'; s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(10.5); s.paragraph_format.space_after=Pt(6); s.paragraph_format.line_spacing=1.18
D.styles['Title'].font.size=Pt(22); D.styles['Heading 1'].font.size=Pt(15)
def p(t): D.add_paragraph(t)
def b(t): D.add_paragraph(t,'List Bullet')
def h(t): D.add_heading(t,1)
def table(headers,rows):
 t=D.add_table(rows=1,cols=len(headers)); t.autofit=False
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=str(x)
 for i,row in enumerate(t.rows):
  for c in row.cells:
   tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF2' if i==0 else 'FFFFFF'); tcPr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
   tcPr.append(borders)
   for para in c.paragraphs:
    para.paragraph_format.space_before=Pt(4); para.paragraph_format.space_after=Pt(4)
    for run in para.runs: run.font.size=Pt(9)
D.add_paragraph('胸部 X 光片多标签分类比赛实施计划','Title')
p('2026年第八届全国高校计算机能力挑战赛 人工智能挑战赛\n计划日期 2026年10月18日    初赛截止 2026年11月13日23:59')
p('先跑通数据读取、训练、检查级评估和提交的完整流程，再通过可比较的实验提高成绩，最后验证模型体积、推理速度和提交格式。本计划将约26天的备赛时间分为六个阶段。')
h('一 明确标签与规则')
p('时间：10月18日至19日')
for x in ['获取官方10类标签的编号与名称映射。','确认预训练权重是否允许使用，以及模型大小和推理延迟的具体测量方式。','分析各类别样本数量、多标签组合与拍摄视角分布，扩大图像完整性检查到全部数据。']: b(x)
p('阶段产出：类别映射表、数据检查报告、明确的训练与提交规范。')
h('二 建立数据读取与评估流程')
p('时间：10月19日至21日')
p('数据层级为患者、检查、图像；一次检查可能包含多张图像。标签与提交结果对应检查，需保留患者编号、检查编号和拍摄视角。')
for x in ['将分号标签转换为长度为10的向量。例如3;5;9表示第3、5、9位为1，其余为0。','按图像读取训练数据；推理时先采用概率平均，合并同次检查的多张图像预测。','使用官方验证集计算Macro-AUC、每类AUC和Macro-F1。后续自行划分数据时必须按患者分组，防止患者交叉。']: b(x)
p('阶段产出：可复用的数据读取、评估与检查级预测汇总流程。')
D.add_page_break()
h('三 跑通第一个基线模型')
p('时间：10月21日至24日')
p('先选中小型模型，例如ResNet-18，尽快建立可运行、可评估、可提交的基线。预训练权重的使用以组委会确认结果为准。')
for x in ['输入先采用320×320，保留图像比例并填充。','输出10个独立概率，推理采用Sigmoid，训练采用BCEWithLogitsLoss。','使用温和的数据增强，避免上下翻转、过大旋转等改变胸片结构的处理。','保存验证集Macro-AUC最高的模型，生成submission.csv并检查格式，尽早尝试一次官方提交。']: b(x)
p('阶段产出：第一份可提交结果、基线成绩、训练耗时和显存占用记录。')
h('四 有控制地提升成绩')
p('时间：10月25日至11月3日')
p('每次只改变一个主要因素，记录配置、结果与收益。优先完成高优先级实验，再决定是否增加复杂度。')
table(['优先级','实验方向','目的'],[
['高','对比320与448输入分辨率','判断细节信息是否有帮助'],['高','对比少量中小型模型','兼顾准确率与推理速度'],['高','调整学习率 轮次 正则化','提高稳定性与泛化能力'],['中','处理类别不均衡','改善少数类别表现'],['中','对比多图平均与视角分组融合','改善检查级预测'],['低','小规模模型融合','在体积与速度允许时提分']])
p('除整体Macro-AUC外，还要观察每类AUC，确认改进是否稳定。当前训练集类别1有622次阳性检查，类别10有6683次，需关注类别不均衡。')
p('阶段产出：实验记录表、表现最好的少量候选模型。')
D.add_page_break()
h('五 优化 F1 并检查部署限制')
p('时间：11月4日至8日')
for x in ['在验证集上为各类别选择阈值，提高Macro-F1。AUC使用连续概率计算，阈值调整主要影响F1，仍应优先优化AUC。','检查模型压缩前总大小不超过500MB。','尽量在规定的T4 16GB环境测量单张图像端到端延迟，目标不超过100ms。','确认训练与推理的预处理一致，并验证检查内多图汇总方法。']: b(x)
p('阶段产出：最终候选模型、类别阈值、推理耗时与模型体积报告。')
h('六 固定最终版本并提交')
p('时间：11月9日至13日')
for x in ['固定模型、预处理与推理程序，避免临近截止时做未经验证的大改动。','校验每个测试检查恰好包含10类。测试集1414次检查，共应有14140条预测记录。','检查概率范围为[0,1]，无缺失、重复或非法值；文件名为submission.csv，编码为UTF-8。','保存模型、配置、依赖版本和完整运行说明，为国赛模型验证与复现做准备。','最迟11月12日完成一次最终候选提交，预留一天处理问题。']: b(x)
p('阶段产出：最终提交文件、完整模型与可复现的推理材料。')
h('执行重点')
p('每天最多提交2次。官方提交应留给有明确改进依据的版本，主要依靠本地验证集做实验。当前先补齐类别映射，再建立读取数据、训练、检查级评估与生成提交文件的最小完整流程。')
p('赛事规则来源：https://www.ncccu.org.cn/index/Paper/case1.html\n本计划依据2026年10月18日查看的赛事信息制定，后续以官方更新为准。')
ROOT = Path(__file__).resolve().parents[2]
D.save(ROOT / 'docs' / '胸部X光片分类比赛实施计划.docx')
